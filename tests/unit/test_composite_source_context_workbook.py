"""Independent OOXML/source/display reconciliation of actual controlled packages."""

import io
import json
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from source_context_fixtures import source_context_package

from app.contracts.render_package import RenderPackage
from app.domain.templates.registry import TemplateRegistry
from app.services.composite_workbook import literal_writer
from app.services.composite_workbook.capacity import preflight_workbook
from app.services.composite_workbook.projection import workbook_tables
from app.services.composite_workbook.rendering import CompositeWorkbookRenderService
from app.services.composite_workbook.source_cells import validate_dataset
from app.services.render_intake import RenderIntakeService
from app.services.render_ports import RenderCompileFailedError


def service() -> CompositeWorkbookRenderService:
    return CompositeWorkbookRenderService(
        RenderIntakeService(TemplateRegistry.load_from_directory(Path("templates/registry")))
    )


def source_scalar(data: Any, pointer: str) -> Any:
    for token in pointer.split("/")[1:]:
        token = token.replace("~1", "/").replace("~0", "~")
        data = data[int(token)] if isinstance(data, list) else data[token]
    return data


def independent_display(column: dict[str, Any], cell: dict[str, Any]) -> str:
    value = cell["canonical_value"]
    if value is None:
        return f"{cell['availability']}: {', '.join(cell['reason_codes'])}"
    if column["value_type"] == "TEXT":
        return str(value)
    number = Decimal(value)
    if column["display_conversion"] == "RATIO_TO_PERCENT_DISPLAY":
        number *= Decimal(100)
    places = column["display_decimal_places"]
    number = number.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    suffix = {"PERCENT": "%", "PERCENTAGE_POINTS": " pp"}.get(column["display_unit"], "")
    if column["value_type"] == "MONEY":
        suffix = f" {column['currency']}"
    return format(number, "f") + suffix


@pytest.mark.parametrize("case,count", [("original", 714), ("corrected", 739)])
def test_complete_independent_workbook_reconciliation(case: str, count: int) -> None:
    wire = source_context_package(case)
    package = RenderPackage.model_validate(wire)
    result = service().render(package)
    if case == "original":
        assert (
            result.artifact_bytes
            == Path("tests/golden/composite-review/v8/expected.xlsx").read_bytes()
        )
    book = openpyxl.load_workbook(io.BytesIO(result.artifact_bytes))
    data = wire["report_data"]
    pinned = "".join(
        str(row[1])
        for sheet in book
        if sheet.title.startswith("PinnedData_")
        for row in list(sheet.values)[1:]
    )
    assert json.loads(pinned) == data
    expected = {
        (t["table_id"], r["row_id"], key): cell
        for t in data["tables"]
        for r in t["rows"]
        for key, cell in r["cells"].items()
    }
    assert len(expected) == count
    for sheet in book:
        if not sheet.title.startswith("CellEvidence_"):
            continue
        for table, row, key, canonical, availability, reasons, pointer in list(sheet.values)[1:]:
            cell = expected.pop((table, row, key))
            source = source_scalar(data, str(pointer))
            exact = (
                None
                if source is None
                else str(source).lower()
                if isinstance(source, bool)
                else str(source)
            )
            assert json.loads(str(canonical)) == exact == cell["canonical_value"]
            assert (availability, json.loads(str(reasons)), pointer) == (
                cell["availability"],
                cell["reason_codes"],
                cell["source_pointer"],
            )
    assert not expected
    policies = {
        (table["table_id"], c["column_id"]): c for table in data["tables"] for c in table["columns"]
    }
    for sheet in book:
        if not sheet.title.startswith("ColumnPolicy_"):
            continue
        for (
            name,
            key,
            label,
            kind,
            unit,
            display,
            conversion,
            places,
            rounding,
            currency,
            scale,
            storage,
        ) in list(sheet.values)[1:]:
            c = policies.pop((name, key))
            assert (
                label,
                kind,
                unit,
                display,
                conversion,
                places,
                rounding,
                currency or "",
                scale,
                storage,
            ) == (
                c["label"],
                c["value_type"],
                c["unit"],
                c["display_unit"],
                c["display_conversion"],
                str(c["display_decimal_places"]),
                c["display_rounding_mode"],
                c["currency"] or "",
                c["scale"],
                "LITERAL_TEXT",
            )
    assert not policies
    identity = {
        row[0]: json.loads(str(row[1])) for row in list(book["ArtifactIdentity_1"].values)[1:]
    }
    assert identity["render_context"] == wire["render_context"]
    assert identity["snapshot_id"] == wire["snapshot_id"]
    assert identity["lineage_refs"] == wire["lineage_refs"]
    for table in data["tables"]:
        actual = list(book[table["table_id"] + "_1"].values)
        assert actual[0] == ("Report row identity", *(c["label"] for c in table["columns"]))
        assert actual[1:] == [
            (
                r["row_id"],
                *(independent_display(c, r["cells"][c["column_id"]]) for c in table["columns"]),
            )
            for r in table["rows"]
        ]
    assert not any(c.data_type == "f" for sheet in book for row in sheet for c in row)
    assert (
        service()
        .render(RenderPackage.model_validate_json(json.dumps(wire, sort_keys=True)))
        .artifact_bytes
        == result.artifact_bytes
    )


def test_physical_capacity_matches_all_evidence_and_refuses_one_below(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    package = RenderPackage.model_validate(source_context_package("corrected"))
    content = validate_dataset(package.report_data)
    digest = (
        TemplateRegistry.load_from_directory(Path("templates/registry"))
        .resolve_for_new_render(package)
        .template_digest
    )
    _, measured = preflight_workbook(package, workbook_tables(package, content, digest))
    book = openpyxl.load_workbook(io.BytesIO(service().render(package).artifact_bytes))
    assert measured.sheets == len(book.sheetnames)
    assert measured.total_cells == sum(sheet.max_row * sheet.max_column for sheet in book)
    monkeypatch.setattr(literal_writer, "MAX_TOTAL_CELLS", measured.total_cells)
    preflight_workbook(package, workbook_tables(package, content, digest))
    monkeypatch.setattr(literal_writer, "MAX_TOTAL_CELLS", measured.total_cells - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit_exceeded"):
        preflight_workbook(package, workbook_tables(package, content, digest))

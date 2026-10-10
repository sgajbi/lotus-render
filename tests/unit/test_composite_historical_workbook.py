"""Independent complete Excel reconciliation of frozen Report r3 packages."""

import io
import json
from pathlib import Path

import openpyxl
import pytest
from historical_fixtures import CASES, historical_package

from app.contracts.composite_historical import CompositeHistoricalContent
from app.contracts.render_package import RenderPackage
from app.domain.templates.registry import TemplateRegistry, TemplateRegistryError
from app.services.composite_workbook import literal_writer
from app.services.composite_workbook.capacity import preflight_workbook
from app.services.composite_workbook.historical_tables import (
    CALCULATION_BOUNDARY,
    validate_historical_table_set,
    validate_historical_tables,
)
from app.services.composite_workbook.projection import workbook_tables
from app.services.composite_workbook.rendering import CompositeWorkbookRenderService
from app.services.composite_workbook.source_cells import validate_dataset
from app.services.render_intake import RenderIntakeService
from app.services.render_ports import RenderCompileFailedError


def _text(value: object) -> str:
    assert isinstance(value, str)
    return value


@pytest.mark.parametrize("case", CASES)
def test_all_physical_evidence_and_sorted_replay(case: str) -> None:
    wire = historical_package(case)
    package = RenderPackage.model_validate(wire)
    registry = TemplateRegistry.load_from_directory(Path("templates/registry"))
    service = CompositeWorkbookRenderService(RenderIntakeService(registry))
    result = service.render(package)
    workbook = openpyxl.load_workbook(io.BytesIO(result.artifact_bytes))
    pinned = "".join(
        _text(row[1])
        for sheet in workbook
        if sheet.title.startswith("PinnedData_")
        for row in list(sheet.values)[1:]
    )
    assert json.loads(pinned) == wire["report_data"]
    expected = {
        (table["table_id"], row["row_id"], key): cell
        for table in wire["report_data"]["tables"]
        for row in table["rows"]
        for key, cell in row["cells"].items()
    }
    for sheet in workbook:
        if not sheet.title.startswith("CellEvidence_"):
            continue
        for table, row, key, canonical, availability, reasons, pointer in list(sheet.values)[1:]:
            cell = expected.pop((table, row, key))
            assert (
                json.loads(_text(canonical)),
                availability,
                json.loads(_text(reasons)),
                pointer,
            ) == (
                cell["canonical_value"],
                cell["availability"],
                cell["reason_codes"],
                cell["source_pointer"],
            )
    assert not expected
    identity = {
        _text(row[0]): json.loads(_text(row[1]))
        for row in list(workbook["ArtifactIdentity_1"].values)[1:]
    }
    assert identity["calculation_boundary"] == CALCULATION_BOUNDARY
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    sorted_package = RenderPackage.model_validate_json(json.dumps(wire, sort_keys=True))
    assert service.render(sorted_package).artifact_bytes == result.artifact_bytes


def test_full_proof_capacity_matches_physical_workbook_and_refuses_one_below(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    package = RenderPackage.model_validate(historical_package())
    registry = TemplateRegistry.load_from_directory(Path("templates/registry"))
    digest = registry.resolve_for_new_render(package).template_digest
    content = validate_dataset(package.report_data)
    _, measured = preflight_workbook(package, workbook_tables(package, content, digest))
    service = CompositeWorkbookRenderService(RenderIntakeService(registry))
    workbook = openpyxl.load_workbook(io.BytesIO(service.render(package).artifact_bytes))
    assert measured.sheets == len(workbook.sheetnames)
    assert measured.total_rows == sum(sheet.max_row - 1 for sheet in workbook)
    assert measured.total_cells == sum(sheet.max_row * sheet.max_column for sheet in workbook)
    monkeypatch.setattr(literal_writer, "MAX_TOTAL_CELLS", measured.total_cells)
    preflight_workbook(package, workbook_tables(package, content, digest))
    monkeypatch.setattr(literal_writer, "MAX_TOTAL_CELLS", measured.total_cells - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit_exceeded"):
        preflight_workbook(package, workbook_tables(package, content, digest))


@pytest.mark.parametrize(
    "change",
    ["missing_row", "duplicate_row", "wrong_pointer", "value", "null", "policy", "disclosure"],
)
def test_provenance_table_refusals(change: str) -> None:
    data = historical_package()["report_data"]
    table = data["tables"][-1]
    cell = table["rows"][0]["cells"]["value"]
    if change == "missing_row":
        table["rows"].pop()
    elif change == "duplicate_row":
        table["rows"][1]["row_id"] = table["rows"][0]["row_id"]
    elif change == "wrong_pointer":
        cell["source_pointer"] = "/report_facts/unavailable"
    elif change == "value":
        cell["canonical_value"] = "changed"
    elif change == "null":
        cell["availability"] = "UNAVAILABLE"
    elif change == "policy":
        table["columns"][2]["display_conversion"] = "RATIO_TO_PERCENT_DISPLAY"
    else:
        data["report_facts"]["disclosures"][-1]["text"] = "Bank verified"
    with pytest.raises(ValueError):
        validate_dataset(data)


@pytest.mark.parametrize("change", ["order", "title", "missing_cell", "root_null"])
def test_table_guards_refuse_schema_valid_projection_changes(change: str) -> None:
    data = historical_package("v1-root-evaluated-only")["report_data"]
    if change == "order":
        data["tables"].reverse()
    elif change == "title":
        data["tables"][-1]["title"] = "False provenance title"
    elif change == "missing_cell":
        data["tables"][-1]["rows"][0]["cells"].pop("evidence_role")
    else:
        data["tables"][-2]["rows"][0]["cells"]["value"]["reason_codes"] = [
            "SOURCE_VALUE_UNAVAILABLE"
        ]
    content = CompositeHistoricalContent.model_validate_json(json.dumps(data))
    with pytest.raises(ValueError, match="composite_historical_"):
        validate_historical_tables(content, data)


def test_layout_cannot_hide_required_proof_tables() -> None:
    layout = json.loads(Path("templates/xlsx/composite-review/v7/layout.json").read_bytes())
    tables = set(layout["required_tables"])
    validate_historical_table_set(layout, tables)
    with pytest.raises(ValueError, match="table_set_invalid"):
        validate_historical_table_set(layout, tables - {"PolicyAdmission"})
    layout["optional_tables"] = ["PolicyAdmission"]
    with pytest.raises(TemplateRegistryError, match="layout_policy_mismatch"):
        validate_historical_table_set(layout, tables)

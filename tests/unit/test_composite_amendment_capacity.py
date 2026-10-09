"""The new consumer's complete lineage fits the existing physical workbook budget."""

import io
from pathlib import Path

import openpyxl
import pytest
from amendment_fixtures import unit_amendment_package

from app.contracts.render_package import RenderPackage
from app.domain.templates.registry import TemplateRegistry
from app.services.composite_workbook import literal_writer
from app.services.composite_workbook.capacity import preflight_workbook
from app.services.composite_workbook.projection import workbook_tables
from app.services.composite_workbook.rendering import CompositeWorkbookRenderService
from app.services.composite_workbook.source_cells import validate_dataset
from app.services.render_intake import RenderIntakeService
from app.services.render_ports import RenderCompileFailedError


def test_full_lineage_preflight_matches_actual_cells_and_refuses_one_below(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    package = RenderPackage.model_validate(unit_amendment_package())
    registry = TemplateRegistry.load_from_directory(Path("templates/registry"))
    digest = registry.resolve_for_new_render(package).template_digest
    content = validate_dataset(package.report_data)
    _, measured = preflight_workbook(package, workbook_tables(package, content, digest))
    service = CompositeWorkbookRenderService(RenderIntakeService(registry))
    artifact = service.render(package).artifact_bytes
    workbook = openpyxl.load_workbook(io.BytesIO(artifact))
    assert measured.sheets == len(workbook.sheetnames)
    assert measured.total_rows == sum(sheet.max_row - 1 for sheet in workbook)
    assert measured.total_cells == sum(sheet.max_row * sheet.max_column for sheet in workbook)
    monkeypatch.setattr(literal_writer, "MAX_TOTAL_CELLS", measured.total_cells)
    preflight_workbook(package, workbook_tables(package, content, digest))
    monkeypatch.setattr(literal_writer, "MAX_TOTAL_CELLS", measured.total_cells - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit_exceeded"):
        preflight_workbook(package, workbook_tables(package, content, digest))

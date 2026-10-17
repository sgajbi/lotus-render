"""V9 projection component controls; full native fixture/workbook proof is separate."""

import io
import json

import openpyxl
import pytest
from source_context_fixtures import source_context_package

from app.contracts.composite_notice import CompositeNoticeContent, NoticeTable
from app.contracts.render_package import RenderPackage
from app.services.composite_workbook import literal_writer
from app.services.composite_workbook.capacity import preflight_workbook
from app.services.composite_workbook.literal_writer import write_literal_workbook
from app.services.composite_workbook.notice_metadata import text_column
from app.services.composite_workbook.projection import workbook_tables
from app.services.render_ports import RenderCompileFailedError


def projection_component() -> tuple[RenderPackage, CompositeNoticeContent]:
    wire = source_context_package()
    wire["template_version"] = "v9"
    data = wire["report_data"]
    data.update(
        {
            "comparison_profile": "component-profile",
            "pair_fingerprint": "component-pair",
            "manifest_digest": "component-manifest",
            "notice_completeness": "PARTIAL",
            "generations": {
                side: {
                    "selection": {"snapshot_id": side},
                    "contract_version": "composite_review.v8",
                    "capture_identity": {"report_revision_id": side},
                    "snapshot_response_bytes_sha256": "component-bytes",
                }
                for side in ("before", "after")
            },
        }
    )
    tables = []
    for name, pointer, value in (
        ("Before_Summary", "/generations/before/payload/report_facts/value", "=1+2"),
        ("After_Summary", "/generations/after/payload/report_facts/value", "@literal"),
        ("NoticeScope", "/notice_facts/scope/0/value", "+literal"),
    ):
        tables.append(
            NoticeTable.model_validate(
                {
                    "table_id": name,
                    "title": name,
                    "columns": [text_column("value")],
                    "rows": [
                        {
                            "row_id": "r/~",
                            "cells": {
                                "value": {
                                    "canonical_value": value,
                                    "availability": "AVAILABLE",
                                    "reason_codes": [],
                                    "source_pointer": pointer,
                                }
                            },
                        }
                    ],
                }
            )
        )
    # Deliberately component-only: this is not a 32-table admitted notice or producer envelope.
    content = CompositeNoticeContent.model_construct(tables=tables)
    return RenderPackage.model_validate(wire), content


def test_v9_evidence_provenance_and_literal_cells_are_complete() -> None:
    package, content = projection_component()
    retained, measured = preflight_workbook(
        package, workbook_tables(package, content, "test-digest")
    )
    artifact = write_literal_workbook(retained)
    assert write_literal_workbook(retained) == artifact
    book = openpyxl.load_workbook(io.BytesIO(artifact))
    evidence = list(book["CellEvidence_1"].values)
    assert len(evidence[0]) == 13 and len(evidence) == 4
    assert evidence[1][7:] == (
        "before",
        "composite_review.v8",
        "Summary",
        "r/~",
        "value",
        "/report_facts/value",
    )
    assert evidence[2][7:] == (
        "after",
        "composite_review.v8",
        "Summary",
        "r/~",
        "value",
        "/report_facts/value",
    )
    assert evidence[3][7:] == ("NOTICE", "", "", "", "", "")
    assert book["Before_Summary_1"]["B2"].value == "=1+2"
    assert book["After_Summary_1"]["B2"].value == "@literal"
    assert book["NoticeScope_1"]["B2"].value == "+literal"
    assert all(
        c.data_type != "f" and c.hyperlink is None for sheet in book for row in sheet for c in row
    )
    assert measured.sheets == len(book.sheetnames)
    assert measured.total_rows == sum(s.max_row - 1 for s in book)
    assert measured.total_cells == sum(s.max_row * s.max_column for s in book)
    assert measured.total_text_bytes == sum(
        len(str(c.value).encode()) for s in book for row in s for c in row if c.value is not None
    )
    identity = {r[0]: json.loads(str(r[1])) for r in list(book["ArtifactIdentity_1"].values)[1:]}
    assert list(identity)[-6:] == [
        "comparison_profile",
        "pair_fingerprint",
        "manifest_digest",
        "before_generation",
        "after_generation",
        "notice_completeness",
    ]
    assert identity["render_context"] == package.render_context
    assert identity["before_generation"] == package.report_data["generations"]["before"]
    assert identity["after_generation"] == package.report_data["generations"]["after"]


@pytest.mark.parametrize(
    "attribute,measured_key",
    [
        ("MAX_TOTAL_ROWS", "total_rows"),
        ("MAX_TOTAL_CELLS", "total_cells"),
        ("MAX_TOTAL_TEXT_BYTES", "total_text_bytes"),
        ("MAX_SHEETS", "sheets"),
    ],
)
def test_all_new_overhead_counts_at_boundary_and_refuses_one_over(
    monkeypatch: pytest.MonkeyPatch,
    attribute: str,
    measured_key: str,
) -> None:
    package, content = projection_component()
    _, measured = preflight_workbook(package, workbook_tables(package, content, "test-digest"))
    bound = getattr(measured, measured_key)
    monkeypatch.setattr(literal_writer, attribute, bound)
    preflight_workbook(package, workbook_tables(package, content, "test-digest"))
    monkeypatch.setattr(literal_writer, attribute, bound - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit_exceeded"):
        preflight_workbook(package, workbook_tables(package, content, "test-digest"))

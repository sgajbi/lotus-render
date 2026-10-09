"""Physical v5 accounting checked against the emitted XLSX, including all evidence."""

import io
import json

import openpyxl
import pytest
from pooled_fixtures import CASES, producer_package
from test_composite_pooled import service, text_value
from test_composite_pooled_refusals import repin

from app.contracts.render_package import RenderPackage
from app.services.composite_workbook import capacity, literal_writer
from app.services.composite_workbook.capacity import preflight_workbook
from app.services.composite_workbook.numeric_display import display_number
from app.services.composite_workbook.projection import workbook_tables
from app.services.composite_workbook.source_cells import validate_dataset
from app.services.render_ports import RenderCompileFailedError

DIGEST = "sha256:d9517f2c8f1b1fdf69294e367f8bee399200be7c59aba8dfed7637d0c0e1fdc6"


@pytest.mark.parametrize("kind", CASES)
def test_v5_capacity_accounts_for_every_emitted_cell_and_partition_header(kind: str) -> None:
    package = RenderPackage.model_validate(producer_package(kind))
    _, measured = preflight_workbook(
        package, workbook_tables(package, validate_dataset(package.report_data), DIGEST)
    )
    book = openpyxl.load_workbook(io.BytesIO(service().render(package).artifact_bytes))
    cells = [text_value(c.value) for sheet in book for row in sheet for c in row]
    assert measured.sheets == len(book.sheetnames)
    assert measured.total_cells == len(cells)
    assert measured.total_text_bytes == sum(len(value.encode("utf-8")) for value in cells)
    assert measured.total_rows == sum(sheet.max_row - 1 for sheet in book)
    book.close()


@pytest.mark.parametrize(
    "bound,field",
    [
        ("MAX_TOTAL_ROWS", "total_rows"),
        ("MAX_TOTAL_CELLS", "total_cells"),
        ("MAX_TOTAL_TEXT_BYTES", "total_text_bytes"),
        ("MAX_SHEETS", "sheets"),
        ("MAX_REQUEST_BODY_BYTES", "request_body_bytes"),
    ],
)
def test_v5_accepts_exact_measured_limit_and_refuses_one_below(
    monkeypatch: pytest.MonkeyPatch, bound: str, field: str
) -> None:
    package = RenderPackage.model_validate(producer_package("corrected"))
    data = validate_dataset(package.report_data)
    _, measured = preflight_workbook(package, workbook_tables(package, data, DIGEST))
    module = capacity if bound == "MAX_REQUEST_BODY_BYTES" else literal_writer
    monkeypatch.setattr(module, bound, getattr(measured, field))
    preflight_workbook(package, workbook_tables(package, data, DIGEST))
    monkeypatch.setattr(module, bound, getattr(measured, field) - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit_exceeded"):
        preflight_workbook(package, workbook_tables(package, data, DIGEST))


def test_exact_large_source_number_and_percent_conversion_remain_bounded() -> None:
    assert display_number("1" * 300 + ".005", 2) == "1" * 300 + ".01"
    assert display_number("0.012345", 4, percent=True) == "1.2345"
    assert display_number("-0.00005", 2, percent=True) == "-0.01"
    with pytest.raises(RenderCompileFailedError, match="resource_limit_exceeded"):
        display_number("1e32767", 2)
    with pytest.raises(ValueError, match="number_invalid"):
        display_number("NaN", 2)


def test_partitioned_source_evidence_and_chunked_identity_recover_complete_input() -> None:
    wire = producer_package()
    wire["report_data"]["source_response"]["outcome"]["diagnostics"]["extended"] = {
        f"key-{index:04d}": str(index) for index in range(1100)
    }
    repin(wire["report_data"])
    wire["render_context"]["archive"]["composite_report_identity"]["selection"] = wire[
        "report_data"
    ]["selection"]
    wire["render_context"]["retained_note"] = "🧾" * 9000
    package = RenderPackage.model_validate(wire)
    _, measured = preflight_workbook(
        package, workbook_tables(package, validate_dataset(package.report_data), DIGEST)
    )
    book = openpyxl.load_workbook(io.BytesIO(service().render(package).artifact_bytes))
    assert "SourceEvidence_2" in book.sheetnames
    identity = {
        text_value(k): text_value(v) for k, v in list(book["ArtifactIdentity_1"].values)[1:]
    }
    descriptor = json.loads(identity["render_context__chunks"])
    restored = "".join(
        json.loads(identity[f"render_context__chunk_{index:06d}"])
        for index in range(descriptor["count"])
    )
    assert json.loads(restored) == wire["render_context"]
    cells = [text_value(c.value) for sheet in book for row in sheet for c in row]
    assert measured.sheets == len(book.sheetnames)
    assert measured.total_cells == len(cells)
    assert measured.total_text_bytes == sum(len(value.encode("utf-8")) for value in cells)
    book.close()

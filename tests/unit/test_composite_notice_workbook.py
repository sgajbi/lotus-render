"""Complete native package through actual writer and independent OOXML decoding."""

import io
import json
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import openpyxl
import pytest
from notice_fixtures import native_json, notice_package

from app.contracts.render_package import RenderPackage
from app.domain.templates.registry import TemplateRegistry
from app.services.composite_workbook import literal_writer
from app.services.composite_workbook.capacity import preflight_workbook
from app.services.composite_workbook.fingerprint import workbook_content_fingerprint
from app.services.composite_workbook.notice_source import validate_notice
from app.services.composite_workbook.projection import workbook_tables
from app.services.composite_workbook.rendering import CompositeWorkbookRenderService
from app.services.render_intake import RenderIntakeService
from app.services.render_ports import RenderCompileFailedError

DIGEST = "sha256:a0a7bf556df61128caedbf6cee69c1ff6a54d365967a66b4ecf5221eaa58c6be"


def service() -> CompositeWorkbookRenderService:
    return CompositeWorkbookRenderService(
        RenderIntakeService(TemplateRegistry.load_from_directory(Path("templates/registry")))
    )


@pytest.fixture(scope="module")
def rendered() -> tuple[Any, Any, Any, Any]:
    package = RenderPackage.model_validate(notice_package())
    content = validate_notice(package.report_data)
    tables, capacity = preflight_workbook(package, workbook_tables(package, content, DIGEST))
    return package, tables, capacity, service().render(package)


def test_complete_native_projection_correspondence(rendered: tuple[Any, Any, Any, Any]) -> None:
    _, tables, _, _ = rendered
    expected = native_json("workbook-projection.json")
    for table, producer in zip(tables, expected, strict=True):
        assert list(table.headers) == producer[0]
        rows = [list(row) for row in table.rows]
        if table.name != "ArtifactIdentity":
            assert rows == producer[1], table.name
            continue
        # Producer planner uses a placeholder digest and a different rounding description/order.
        # Receiver retains its established identity fields and order, appending only six fields.
        actual_identity, planned_identity = dict(rows), dict(producer[1])
        assert set(actual_identity) == set(planned_identity)
        for key in actual_identity.keys() - {"template_digest", "display_rounding"}:
            assert actual_identity[key] == planned_identity[key], key
        assert rows[-6:] == producer[1][-6:]
        assert actual_identity["template_digest"] == json.dumps(DIGEST)
        assert json.loads(actual_identity["display_rounding"]) == (
            "Declared column decimal places; HALF_UP; ratio times 100 once"
        )


def test_all_physical_cells_partitions_and_pinned_roundtrip(
    rendered: tuple[Any, Any, Any, Any],
) -> None:
    package, tables, capacity, result = rendered
    book = openpyxl.load_workbook(io.BytesIO(result.artifact_bytes))
    expected = {}
    for table in tables:
        rows = list(table.rows)
        for offset in range(0, len(rows), 1000):
            name = f"{table.name}_{offset // 1000 + 1}"
            expected[name] = [table.headers, *rows[offset : offset + 1000]]
    assert book.sheetnames == list(expected)
    assert len({name.casefold() for name in book.sheetnames}) == capacity.sheets == 47
    cells = text_bytes = data_rows = 0
    for sheet in book:
        actual = [tuple("" if c.value is None else c.value for c in row) for row in sheet]
        assert actual == expected[sheet.title], sheet.title
        assert sheet.max_row <= 1001
        data_rows += sheet.max_row - 1
        for row in sheet:
            cells += len(row)
            for cell in row:
                assert cell.data_type != "f" and cell.hyperlink is None
                value = "" if cell.value is None else cell.value
                assert isinstance(value, str)
                text_bytes += len(value.encode("utf-8"))
    assert (data_rows, cells, text_bytes) == (12777, 150148, 7746370)
    assert (capacity.total_rows, capacity.total_cells, capacity.total_text_bytes) == (
        data_rows,
        cells,
        text_bytes,
    )
    pinned = "".join(
        str(row[1])
        for sheet in book
        if sheet.title.startswith("PinnedData_")
        for row in list(sheet.values)[1:]
    )
    assert json.loads(pinned) == package.report_data
    with ZipFile(io.BytesIO(result.artifact_bytes)) as archive:
        assert not any("externalLink" in name for name in archive.namelist())


def independent_display(column: dict[str, Any], cell: dict[str, Any]) -> str:
    value = cell["canonical_value"]
    if value is None:
        return f"{cell['availability']}: {', '.join(cell['reason_codes'])}"
    if column["value_type"] == "TEXT":
        return str(value)
    number = Decimal(value)
    if column["display_conversion"] == "RATIO_TO_PERCENT_DISPLAY":
        number *= Decimal(100)
    number = number.quantize(Decimal(1).scaleb(-column["display_decimal_places"]), ROUND_HALF_UP)
    suffix = {"PERCENT": "%", "PERCENTAGE_POINTS": " pp"}.get(column["display_unit"], "")
    if column["value_type"] == "MONEY":
        suffix = f" {column['currency']}"
    return format(number, "f") + suffix


def test_every_visible_cell_uses_exact_source_policy(rendered: tuple[Any, Any, Any, Any]) -> None:
    package, _, _, result = rendered
    book = openpyxl.load_workbook(io.BytesIO(result.artifact_bytes))
    count = 0
    for table in package.report_data["tables"]:
        for index, row in enumerate(table["rows"]):
            sheet = book[f"{table['table_id']}_{index // 1000 + 1}"]
            actual = tuple(
                sheet.iter_rows(
                    min_row=index % 1000 + 2, max_row=index % 1000 + 2, values_only=True
                )
            )[0]
            expected = (
                row["row_id"],
                *(independent_display(c, row["cells"][c["column_id"]]) for c in table["columns"]),
            )
            assert actual == expected
            count += len(table["columns"])
    assert count == sum(len(t["rows"]) * len(t["columns"]) for t in package.report_data["tables"])


def test_actual_output_fit_determinism_and_output_cap(
    rendered: tuple[Any, Any, Any, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    package, _, capacity, result = rendered
    assert capacity.request_body_bytes == 4824578 < 8388608
    assert len(result.artifact_bytes) == 890817 < 16777216
    assert workbook_content_fingerprint(result.artifact_bytes) == (
        "89bdd847cd33ebceebeb8127039e79909627828f72a05ef7f5c8b5c22ccdfab1"
    )
    monkeypatch.setattr(literal_writer, "MAX_OUTPUT_BYTES", len(result.artifact_bytes))
    # Test writer directly: registry correctly refuses monkeypatched policy mismatches.
    content = validate_notice(package.report_data)
    layout = json.loads(Path("templates/xlsx/composite-review/v9/layout.json").read_bytes())
    colors = (layout["header_background"], layout["header_foreground"])
    exact = literal_writer.write_literal_workbook(
        workbook_tables(package, content, DIGEST), header_colors=colors
    )
    assert len(exact) == len(result.artifact_bytes)
    assert workbook_content_fingerprint(exact) == workbook_content_fingerprint(
        result.artifact_bytes
    )
    monkeypatch.setattr(literal_writer, "MAX_OUTPUT_BYTES", len(exact) - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit"):
        literal_writer.write_literal_workbook(
            workbook_tables(package, content, DIGEST), header_colors=colors
        )


def test_full_request_boundary_and_complete_producer_budget(
    rendered: tuple[Any, Any, Any, Any],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.services.composite_workbook import capacity as limits

    package, tables, actual, _ = rendered
    producer = native_json("producer-budget-proof.json")
    # Producer's raw budget proof records its own identity planning differences.
    for key in ("request_body_bytes", "sheets", "total_rows", "total_cells"):
        assert getattr(actual, key) == producer["measured"][key]
    assert actual.total_text_bytes + 7 == producer["measured"]["total_text_bytes"]
    monkeypatch.setattr(limits, "MAX_REQUEST_BODY_BYTES", actual.request_body_bytes)
    preflight_workbook(package, tables)
    monkeypatch.setattr(limits, "MAX_REQUEST_BODY_BYTES", actual.request_body_bytes - 1)
    with pytest.raises(RenderCompileFailedError, match="resource_limit"):
        preflight_workbook(package, tables)

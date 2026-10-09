"""The genuine controlled six-year package fits the registered bounded writer on every CI host."""

import gzip
import hashlib
import io
import json
from pathlib import Path

import openpyxl
import pytest

from app.contracts.render_package import RenderPackage
from app.core.settings import Settings
from app.domain.templates.registry import TemplateRegistry
from app.services.composite_workbook.rendering import CompositeWorkbookRenderService
from app.services.render_intake import RenderIntakeService


def _text(value: object) -> str:
    assert isinstance(value, str)
    return value


@pytest.mark.parametrize(
    "fixture,sha256,wire_bytes,sheets,rows,cells,semantic_cells,context_bytes",
    [
        (
            "v1/six-year-render-package.json.gz",
            "a66e8cb356ccb1dba335e874494cb584e827b086316f3684a26114da50d1355a",
            6_893_642,
            43,
            27_939,
            201_586,
            24_604,
            60_148,
        ),
        (
            "v2/original-six-year-render-package.json.gz",
            "2285991971739d839c9ca1cf7f6a1e439fac571266d8ed27b498a3ab1c255d77",
            7_565_462,
            44,
            28_020,
            202_071,
            24_623,
            None,
        ),
        (
            "v2/financial-correction-six-year-render-package.json.gz",
            "5c029414d124f8d58ee566632d5bd5373abace93f5635aff28f3d83c365d5f84",
            7_566_158,
            44,
            28_020,
            202_071,
            24_623,
            None,
        ),
    ],
)
def test_actual_six_year_package_preserves_all_evidence_pins_and_large_context(
    fixture: str,
    sha256: str,
    wire_bytes: int,
    sheets: int,
    rows: int,
    cells: int,
    semantic_cells: int,
    context_bytes: int | None,
) -> None:
    raw = gzip.decompress((Path("tests/golden/composite-review") / fixture).read_bytes())
    assert hashlib.sha256(raw).hexdigest() == sha256
    payload = json.loads(raw)
    wire = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode()
    assert len(wire) == wire_bytes < Settings().max_request_body_bytes
    package = RenderPackage.model_validate(payload)
    service = CompositeWorkbookRenderService(
        RenderIntakeService(TemplateRegistry.load_from_directory(Path("templates/registry")))
    )
    result = service.render(package)
    assert len(result.artifact_bytes) < 16_777_216
    workbook = openpyxl.load_workbook(io.BytesIO(result.artifact_bytes), read_only=True)
    assert len(workbook.sheetnames) == sheets
    assert sum(sheet.max_row - 1 for sheet in workbook) == rows
    assert sum(sheet.max_row * sheet.max_column for sheet in workbook) == cells
    pinned = "".join(
        _text(row[1])
        for sheet in workbook
        if sheet.title.startswith("PinnedData_")
        for row in list(sheet.values)[1:]
    )
    assert json.loads(pinned) == payload["report_data"]
    expected = {
        (table["table_id"], row["row_id"], column["column_id"]): row["cells"][column["column_id"]]
        for table in payload["report_data"]["tables"]
        for row in table["rows"]
        for column in table["columns"]
    }
    actual = [
        row
        for sheet in workbook
        if sheet.title.startswith("CellEvidence_")
        for row in list(sheet.values)[1:]
    ]
    assert len(actual) == len(expected) == semantic_cells
    assert len({row[:3] for row in actual}) == len(expected)
    for table, row, column, canonical, availability, reasons, pointer in actual:
        cell = expected[(table, row, column)]
        assert json.loads(_text(canonical)) == cell["canonical_value"]
        assert availability == cell["availability"]
        assert json.loads(_text(reasons)) == cell["reason_codes"]
        assert pointer == cell["source_pointer"]
    identity = list(workbook["ArtifactIdentity_1"].values)[1:]
    position = next(i for i, row in enumerate(identity) if row[0] == "render_context__chunks")
    descriptor = json.loads(_text(identity[position][1]))
    fragments = identity[position + 1 :]
    expected_context = json.dumps(payload["render_context"], ensure_ascii=True, sort_keys=True)
    count = (len(expected_context) + 15_999) // 16_000
    assert descriptor["count"] == len(fragments) == count
    assert [row[0] for row in fragments] == [
        f"render_context__chunk_{index:06d}" for index in range(count)
    ]
    canonical = "".join(json.loads(_text(row[1])) for row in fragments)
    assert canonical == expected_context
    assert len(canonical.encode()) == descriptor["utf8_bytes"]
    if context_bytes is not None:
        assert descriptor["utf8_bytes"] == context_bytes
    assert hashlib.sha256(canonical.encode()).hexdigest() == descriptor["sha256"]
    assert json.loads(canonical) == payload["render_context"]
    workbook.close()

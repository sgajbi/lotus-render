"""Independent Excel parser verifies literal safety, deterministic bytes and full partitioning."""

import io
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

import openpyxl
import pytest

from app.services.composite_workbook import literal_writer as writer
from app.services.composite_workbook.literal_writer import LiteralTable, write_literal_workbook
from app.services.render_ports import RenderCompileFailedError, RenderEngineTimeoutError


def _table(values: tuple[str, ...]) -> LiteralTable:
    return LiteralTable("Summary", ("Exact evidence",), ((value,) for value in values))


def test_text_is_literal_in_excel_including_formula_url_and_large_number() -> None:
    values = (
        '=HYPERLINK("https://invalid.example","click")',
        "+SUM(A1:A2)",
        "-1+2",
        "@SUM(1,2)",
        "https://invalid.example/path",
        "000000000000000000000001",
        "999999999999999999999999.1234567890123456789",
        "_x000D_ literal",
        "汉字 🚀",
    )
    artifact = write_literal_workbook([_table(values)])
    workbook = openpyxl.load_workbook(io.BytesIO(artifact))
    sheet = workbook["Summary_1"]
    assert [cell.value for row in sheet.iter_rows(min_row=2) for cell in row] == list(values)
    assert all(cell.data_type == "s" and cell.hyperlink is None for row in sheet for cell in row)
    with ZipFile(io.BytesIO(artifact)) as package:
        xml = b"".join(package.read(name) for name in package.namelist() if name.endswith(".xml"))
        assert b"<f>" not in xml
        assert b"<hyperlink" not in xml
        assert not any("vba" in name.lower() for name in package.namelist())
    workbook.close()


def test_repeated_render_has_identical_bytes_under_fixed_metadata_policy() -> None:
    first = write_literal_workbook([_table(("0.010000000000", "0.030200000000"))])
    second = write_literal_workbook([_table(("0.010000000000", "0.030200000000"))])
    assert first == second


def test_partition_preserves_order_and_repeats_header_without_truncation() -> None:
    count = writer.MAX_ROWS_PER_SHEET + 1
    values = tuple(str(index) for index in range(count))
    artifact = write_literal_workbook([_table(values)])
    workbook = openpyxl.load_workbook(io.BytesIO(artifact), read_only=True)
    assert workbook.sheetnames == ["Summary_1", "Summary_2"]
    actual: list[object] = []
    for sheet in workbook:
        rows = list(sheet.values)
        assert rows[0] == ("Exact evidence",)
        actual.extend(row[0] for row in rows[1:])
    assert actual == list(values)
    workbook.close()


@pytest.mark.parametrize("value", ["\x00", "\x01", "\x0b", "\x1f"])
def test_invalid_xml_text_is_refused(value: str) -> None:
    with pytest.raises(ValueError, match="composite_workbook_text_invalid"):
        write_literal_workbook([_table((value,))])


def test_excel_utf16_text_boundary_accepts_exact_limit_and_refuses_one_more() -> None:
    writer.validate_literal("a" * writer.MAX_CELL_UTF16_UNITS)
    with pytest.raises(RenderCompileFailedError):
        writer.validate_literal("a" * (writer.MAX_CELL_UTF16_UNITS + 1))
    with pytest.raises(RenderCompileFailedError):
        writer.validate_literal("🚀" * (writer.MAX_CELL_UTF16_UNITS // 2 + 1))


@pytest.mark.parametrize(
    "limit,maximum",
    [
        ("MAX_TOTAL_ROWS", 1),
        ("MAX_TOTAL_CELLS", 2),
        ("MAX_TOTAL_TEXT_BYTES", 1),
        ("MAX_SHEETS", 1),
        ("MAX_OUTPUT_BYTES", 1),
    ],
)
def test_resource_guards_refuse_without_partial_artifact(
    monkeypatch: pytest.MonkeyPatch, limit: str, maximum: int
) -> None:
    # The full valid writer path passed above. Lower limits exercise each refusal path.
    monkeypatch.setattr(writer, limit, maximum)
    tables = [_table(("one", "two")), LiteralTable("Lineage", ("id",), [("three",)])]
    with pytest.raises(RenderCompileFailedError) as refusal:
        write_literal_workbook(tables)
    assert refusal.value.failure_category.value == "resource_limit_exceeded"


def test_excess_columns_and_incomplete_rows_refuse() -> None:
    for table in (
        LiteralTable("Summary", ("column",) * (writer.MAX_COLUMNS + 1), []),
        LiteralTable("Summary", ("one", "two"), [("missing",)]),
    ):
        with pytest.raises(RenderCompileFailedError):
            write_literal_workbook([table])


def test_no_tables_never_produces_a_plausible_empty_workbook() -> None:
    with pytest.raises(ValueError, match="composite_workbook_tables_missing"):
        write_literal_workbook([])


def test_expired_deadline_stops_workbook_and_cleans_temporary_files(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    created: list[Path] = []

    def temporary_directory(*, prefix: str) -> TemporaryDirectory[str]:
        directory = TemporaryDirectory(prefix=prefix, dir=tmp_path)
        created.append(Path(directory.name))
        return directory

    monkeypatch.setattr(writer, "TemporaryDirectory", temporary_directory)
    with pytest.raises(RenderEngineTimeoutError, match="composite_workbook_timeout"):
        write_literal_workbook([_table(("value",))], timeout_seconds=0)
    assert created and all(not path.exists() for path in created)

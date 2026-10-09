"""Independent Excel parsing of registered v4 rendering from frozen Report unit packages."""

import io
import json
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from eligibility_fixtures import producer_package

from app.contracts.render_package import RenderPackage
from app.domain.templates.registry import TemplateRegistry
from app.services.composite_workbook.rendering import CompositeWorkbookRenderService
from app.services.render_intake import RenderIntakeService


def _text(value: object) -> str:
    assert isinstance(value, str)
    return value


def _display(column: dict[str, Any], cell: dict[str, Any]) -> str:
    value = cell["canonical_value"]
    if value is None:
        return str(cell["availability"]) + ": " + ", ".join(cell["reason_codes"])
    if column["value_type"] in {"TEXT", "BOOLEAN"}:
        return str(value)
    places = column["display_decimal_places"]
    number = Decimal(value).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    return format(number, "f") + (
        " " + column["currency"] if column["value_type"] == "MONEY" else ""
    )


@pytest.mark.parametrize(
    "kind,reasons",
    [
        ("evaluated_only", 7),
        ("published", 4),
        ("actual_v1", 4),
        ("actual_v2", 4),
    ],
)
def test_registered_workbook_preserves_every_cell_policy_and_retained_dataset(
    kind: str, reasons: int
) -> None:
    wire = producer_package(kind)
    package = RenderPackage.model_validate(wire)
    service = CompositeWorkbookRenderService(
        RenderIntakeService(TemplateRegistry.load_from_directory(Path("templates/registry")))
    )
    result = service.render(package)
    assert service.render(package).artifact_bytes == result.artifact_bytes
    workbook = openpyxl.load_workbook(io.BytesIO(result.artifact_bytes))
    assert len(workbook.sheetnames) == 12
    assert not any(name.startswith("CanonicalData") for name in workbook.sheetnames)
    data = wire["report_data"]
    expected = {}
    for table in data["tables"]:
        visible = list(workbook[table["table_id"] + "_1"].values)
        assert visible[0] == (
            "Report row identity",
            *(column["label"] for column in table["columns"]),
        )
        assert visible[1:] == [
            (
                row["row_id"],
                *(
                    _display(column, row["cells"][column["column_id"]])
                    for column in table["columns"]
                ),
            )
            for row in table["rows"]
        ]
        for row in table["rows"]:
            for column in table["columns"]:
                expected[(table["table_id"], row["row_id"], column["column_id"])] = row["cells"][
                    column["column_id"]
                ]
    evidence = list(workbook["CellEvidence_1"].values)[1:]
    assert len(evidence) == len(expected)
    for table_id, row_id, column_id, canonical, availability, reason_codes, pointer in evidence:
        cell = expected.pop((table_id, row_id, column_id))
        assert (
            json.loads(_text(canonical)),
            availability,
            json.loads(_text(reason_codes)),
            pointer,
        ) == (
            cell["canonical_value"],
            cell["availability"],
            cell["reason_codes"],
            cell["source_pointer"],
        )
    assert not expected
    pinned = "".join(_text(row[1]) for row in list(workbook["PinnedData_1"].values)[1:])
    assert json.loads(pinned) == data
    identity = {
        key: json.loads(_text(value))
        for key, value in list(workbook["ArtifactIdentity_1"].values)[1:]
    }
    assert identity["render_context"] == wire["render_context"]
    assert (
        identity["display_rounding"]
        == "Declared column decimal places; HALF_UP; source ratios remain ratios"
    )
    assert workbook["EligibilityReasons_1"].max_row - 1 == reasons
    policies = list(workbook["ColumnPolicy_1"].values)[1:]
    assert policies == [
        (
            table["table_id"],
            column["column_id"],
            column["label"],
            column["value_type"],
            column["unit"],
            column["display_unit"],
            column["display_conversion"],
            str(column["display_decimal_places"]),
            column["display_rounding_mode"],
            column["currency"] or "",
            column["scale"],
            "LITERAL_TEXT",
        )
        for table in data["tables"]
        for column in table["columns"]
    ]
    assert all(
        cell.data_type == "s" and cell.hyperlink is None
        for sheet in workbook
        for row in sheet
        for cell in row
    )
    workbook.close()
